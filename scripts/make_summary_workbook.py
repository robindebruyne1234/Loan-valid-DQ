"""Build deliverables/Berka_Summary_Workbook.xlsx: six tabs, one story.

1 Defects · 2 Model & Explanation · 3 Semantic solution · 4 Questions & Answers
· 5 Assumptions · 6 Design decisions

Nothing here is typed from memory where the repository already holds the
fact: defects come from ``deliverables/defects.csv``, row counts from the
built database, the generated SQL and answers from ``dq/results`` and
``llm_logs``, assumptions from ``ASSUMPTIONS.md`` and decisions from
``DECISIONS.md``. Only the pipeline narrative (tab 3) is written here.

Read-only against the database. Usage: python scripts/make_summary_workbook.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src import build_model, config, db  # noqa: E402
from src import make_workbook as mw  # noqa: E402

TARGET = config.DELIVERABLES_DIR / "Berka_Summary_Workbook.xlsx"
GREEN = mw.PatternFill("solid", fgColor="E2EFDA")
AMBER = mw.PatternFill("solid", fgColor="FFF2CC")


# --------------------------------------------------------------------------
# Tab 3: the pipeline narrative
# --------------------------------------------------------------------------

#: (step, stage, command, module, what happens, what to inspect, why)
PIPELINE: list[tuple[str, ...]] = [
    ("1", "Ingest (Bronze)",
     "make ingest\nscripts\\windows.ps1 ingest",
     "src/ingest.py",
     "Loads the 8 source files (1,079,680 rows) into schema raw with COPY, every column as text, "
     "unchanged. Adds _source_file and _source_row, and records a SHA-256, byte size, rows "
     "loaded and rows expected per file in raw.load_audit. trans.asc.gz is read compressed.",
     "raw.* (text only) and raw.load_audit (counts_match = true for all 8 files).",
     "Proof of what arrived. Nothing is typed, cleaned or dropped here, so every later number "
     "can be traced back to a source row."),
    ("2", "Profile and stage (Silver)",
     "make profile\nscripts\\windows.ps1 profile",
     "src/profile_data.py, sql/10_stg.sql",
     "Builds schema stg: types the columns, decodes the Czech codes to labels (28 codes), "
     "derives gender and birth date from birth_number (month + 50 = female), resolves the sign "
     "of every transaction once (signed_amount: PRIJEM +, VYDAJ and VYBER -), turns placeholders "
     "into NULL. Never filters a row and never repairs a value: the original text survives in "
     "*_raw columns with a _conversion_notes entry. Writes dq/results/profile.md and .json.",
     "stg.* (row counts equal raw), profile.md, docs/CLEANING_LEDGER.md (21 cleaning steps).",
     "Typing and decoding are done once, in one place, so no consumer re-implements them. "
     "'Staging casts, it never filters' keeps the rule layer able to report on the source."),
    ("3", "Data quality",
     "make dq\nscripts\\windows.ps1 dq",
     "src/dq_runner.py, dq/dq_rules.yaml",
     "Runs 124 rules (50 block, 23 warn, 51 info) over completeness, uniqueness, referential "
     "integrity, validity and cross-table consistency. block -> row goes to dq.quarantine and "
     "is excluded from core. warn -> row kept, flag added to _dq_flags. info -> logged. Also "
     "reconciles raw -> stg on row counts and money, and runs the same rules against a copy of "
     "staging with 20 seeded defects (schema inject) to prove each rule fires.",
     "dq.results, dq.quarantine (1 row), dq.flags (440 rows), dq.reconciliation, "
     "dq/results/dq_report.md, deliverables/defects.csv.",
     "A severity must have a mechanical consequence, otherwise it is a comment."),
    ("4", "Model (Gold)",
     "make model\nscripts\\windows.ps1 model",
     "src/build_model.py, sql/30_core.sql",
     "Builds schema core as a star schema with a bridge: dim_date, dim_district, dim_client, "
     "dim_account, bridge_disposition (owner or disponent), dim_card, fact_loan "
     "(is_problematic = status B or D), fact_transaction (1,056,319 rows = stg minus 1 "
     "quarantined), fact_account_month_end_balance (month-end balance: the closing posting is "
     "the end of the day's running-balance chain, D028, with carry-forward) and "
     "ref_codes. Adds curated views v_account_owner, v_loan_enriched, v_client_region. Adds "
     "opaque client_ref and account_ref. Then re-applies the ai views (step 7).",
     "core.* with row counts reconciled: core = stg - quarantine on rows and on money.",
     "Four grains (client, account, loan, posting) appear in the ten questions. One flat table "
     "would force count(distinct) everywhere; a star with a bridge keeps counts natural."),
    ("5", "Cleaning ledger and reference answers",
     "make ledger\nmake reference\nscripts\\windows.ps1 ledger\nscripts\\windows.ps1 reference",
     "src/cleaning_ledger.py, src/run_reference.py, sql/60_reference_queries.sql",
     "The ledger measures each cleaning step on the live database and writes the evidence. "
     "The reference step runs 10 hand-written SQL queries (never shown to the LLM) and stores "
     "the ground truth in dq/results/reference_results.json.",
     "docs/CLEANING_LEDGER.md, dq/results/reference_results.json.",
     "The LLM is graded against an answer that was written independently of it."),
    ("6", "Semantic layer",
     "make semantic\nscripts\\windows.ps1 semantic",
     "semantic/semantic_layer.yaml, src/semantic_context.py",
     "The YAML describes each ai entity (grain, key, columns, join paths), the certified "
     "metrics, a glossary that says what each term is NOT, and the data-quality caveats with "
     "row counts. The build validates it against information_schema: a described column that "
     "does not exist, an existing column nobody described, or a grain that is not really one "
     "row per key fails the build. It is then rendered into semantic/context_pack.md "
     "(about 36,900 characters), which is all the LLM ever sees.",
     "semantic/context_pack.md.",
     "A semantic layer that has drifted from the schema is worse than none: the model cannot "
     "check it."),
    ("7", "Security: authorisation (who may see what)",
     "make entitlements\nscripts\\windows.ps1 entitlements",
     "src/build_security.py, sql/40_ai_views.sql, sql/50_sec_rls.sql",
     "Creates the ai views (all security_invoker, so row level security runs as the caller), "
     "the consumer roles, and 7 row level security policies on the client-bearing core tables "
     "(identity = session_user, which no statement can forge). Access comes from "
     "sec.v_access_grant: ownership, disposition (a disponent reaches the account, not the "
     "person), portfolio (relationship managers, view only) and delegation (time-boxed, the "
     "only stored grant). Onboarding a user or portfolio is inserts, not code.",
     "ai.* views, sec.* tables, sec.v_who_can_reach_client ('who can see this client, and why').",
     "The filter sits on the base tables in the engine. The LLM only emits text, and the "
     "predicate is not in the text, so rewriting the query cannot remove it."),
    ("8", "Security: authentication (two factors)",
     "(part of make entitlements)",
     "sql/51_sec_auth.sql, src/auth.py",
     "Factor 1 is the database password. Factor 2 is a time-based one-time code (RFC 6238 TOTP, "
     "implemented in SQL). A password alone resolves every entitlement to nothing: 0 rows. "
     "sec.verify_second_factor(code) takes only a code, never a principal (it uses "
     "session_user). A code works once, five failures lock the factor for 15 minutes, and a "
     "successful login opens a session bound to that database connection for 30 minutes.",
     "sec.auth_factor, sec.auth_session; tests/test_auth.py (39 tests).",
     "Two-factor is enforced where the data is, so no application can skip it."),
    ("9", "Demo setup",
     "make demo\nmake demo-login\nmake web-open\nscripts\\windows.ps1 demo / demo-login / web-open",
     "src/worked_example.py, src/demo_console.py, src/webapp.py, web/",
     "Onboards demo logins and enrols their second factor: relationship managers (rm_north, "
     "rm_south, rm_unassigned), a risk user (cro_global), an administrator (admin_demo) and "
     "three customers (client_3763, client_3764, client_11409). 'demo' is the worked example "
     "for one client. 'demo-login' is a 7-act console walkthrough: password only -> 0 rows, "
     "password + code -> the user's slice, switching between logins. 'web-open' serves the "
     "7-step browser walkthrough on http://localhost:8080 (bronze, silver clean, silver grade, "
     "gold, access, login, ask) with the design choice shown on each page.",
     "docs/WORKED_EXAMPLE.md; the console; the web page.",
     "The demo prints the one-time codes on screen. That is the point of a demo and the "
     "opposite of security; every persona shares the password demo_pw. Acceptable only "
     "because Berka is public 1999 data."),
    ("10", "Asking questions",
     "make questions\nmake evaluate\nscripts\\windows.ps1 questions / evaluate\n(web: step 7 Ask)",
     "src/run_questions.py, src/llm_client.py, src/sql_guard.py, src/sql_ir.py, src/evaluate.py",
     "For each question: (1) the context pack plus the question is sent to the LLM, which has "
     "no tools, no database and no files and must answer with one JSON object: sql, ir (the "
     "certified metric, entity, counting grain, half-open periods, grouping, filters it "
     "intends), reasoning, assumptions, caveats. (2) sqlglot parses the SQL: exactly one "
     "read-only statement, only the ai schema, a row cap added, a statement timeout set. "
     "(3) src/sql_ir.py checks the IR against the semantic layer and the SQL against the IR: "
     "every column exists, the period filter selects exactly the IR's days, joins are declared "
     "and cannot repeat the counted entity, grouping and filters are the ones declared (D029). "
     "(4) Only then does it log in as the chosen database user (both factors), have PostgreSQL "
     "plan the query, run it under row level security, and recompute the metric from the IR; "
     "the two must agree before any row is shown. (5) Any refusal allows one retry with the "
     "reason (maximum 2 attempts). (6) Every call is logged in llm_logs/. (7) 'evaluate' "
     "compares each verified answer value-by-value with the reference (tolerance 0.011). "
     "'make extra' does the same for the additional gender and quarter questions.",
     "dq/results/llm_results.json, dq/results/evaluation.json, llm_logs/Q*_attempt*.json; "
     "dq/results/*_extra.json and llm_logs/extra/ for the additional questions.",
     "The LLM is treated as an untrusted query author: the guard checks the shape, the IR "
     "verifier checks the meaning, the policy decides the rows. A guard pass alone proves "
     "nothing about correctness."),
]

#: (topic, detail) for the Ollama block
OLLAMA: list[tuple[str, str]] = [
    ("What Ollama is here",
     "A local server (http://localhost:11434) that runs an open-weight model on your own "
     "machine. In this project it is an alternative provider behind the same interface as the "
     "Claude CLI and the Anthropic API (src/llm_client.py, call_ollama). Free, no API key, and "
     "no data leaves the machine."),
    ("Status in this repository (read this first)",
     "The ten recorded answers on tab 4, and the four additional ones below them, were "
     "generated by Claude Sonnet 5.5 through the Claude CLI, not by Ollama. The code path for "
     "Ollama exists and is covered by the provider switch, but the note in src/llm_client.py "
     "and in DECISIONS.md D005 says it has never executed in the build environment, because "
     "the model download hosts were not reachable there; no Ollama server was reachable for "
     "the 2026-10-28 re-evaluation either. No Ollama result is claimed anywhere in this "
     "workbook."),
    ("Install (Windows)",
     "1. Install Ollama from ollama.com (Windows installer; it starts a background service).\n"
     "2. PowerShell: ollama pull qwen2.5-coder:7b\n"
     "   (lighter: qwen2.5-coder:1.5b, about 2 GB RAM. Heavier models write better SQL but are "
     "slower.)\n"
     "3. Check: ollama list   and   Invoke-RestMethod http://localhost:11434/api/tags"),
    ("Point the project at it (PowerShell, native Windows)",
     "$env:LLM_PROVIDER = \"ollama\"\n"
     "$env:LLM_MODEL = \"qwen2.5-coder:7b\"\n"
     ".venv\\Scripts\\python.exe -m src.run_questions --out dq\\results\\llm_results_ollama.json\n"
     ".venv\\Scripts\\python.exe -m src.evaluate --llm dq\\results\\llm_results_ollama.json "
     "--out dq\\results\\evaluation_ollama.json\n"
     "Additional questions: add --suite extra and use llm_results_extra_ollama.json / "
     "evaluation_extra_ollama.json.\n"
     "Use separate files like this: the default ones hold the Claude run, and the runner only "
     "merges results from the same provider and model. Call logs go to "
     "llm_logs\\llm_results_ollama\\. Optional: $env:OLLAMA_HOST (default "
     "http://localhost:11434), $env:LLM_TIMEOUT_SECONDS (default 300). Linux/macOS/WSL: export "
     "LLM_PROVIDER=ollama LLM_MODEL=qwen2.5-coder:7b and the same python -m commands."),
    ("Asking from the web page",
     "Start the site in the same window after setting the two variables "
     "(scripts\\windows.ps1 web-open). Step 7 'Ask' then calls the local model. If Ollama is not "
     "running, or the model is not pulled, the page shows NOT EXECUTED rather than inventing "
     "an answer."),
    ("What is sent",
     "POST /api/generate with the system prompt, the rendered context pack plus the question, "
     "stream = false and temperature 0.0 (src/config.py LLM_TEMPERATURE). The reply is parsed "
     "as one JSON object {sql, ir, reasoning, assumptions, caveats}."),
    ("What does not change with the provider",
     "The semantic layer, the SQL guard, the IR verifier, the row cap and timeout, the retry "
     "rule (max 2 attempts), the login (password + one-time code) and row level security are "
     "identical. A weaker model can therefore write a worse query, but the verifier refuses it "
     "rather than showing a number it cannot vouch for, and it cannot see rows its logged-in "
     "user is not entitled to."),
    ("What to expect",
     "Smaller models will get some of the ten wrong or return malformed JSON. That is worth "
     "showing: 'evaluate' reports exactly which questions match, are partial or mismatch the "
     "reference. Compare against the recorded Claude run on tab 4. Expectation, not a "
     "measurement: this workbook has no Ollama numbers."),
    ("Fail-safe behaviour",
     "No provider or an unreachable one -> every cell reads NOT EXECUTED and the command exits "
     "non-zero. A fabricated AI result would be worse than none."),
]


def sheet_semantic(book: Workbook) -> None:
    """Write tab 3: the pipeline from ingest to asking questions, with Ollama.

    :param book: Target workbook.
    :returns: ``None``.
    """
    sheet = book.create_sheet("3 Semantic solution")
    mw.set_widths(sheet, [6, 24, 40, 34, 78, 44, 52])
    sheet.cell(row=1, column=1,
               value="Semantic solution: ingest -> profile -> model -> security -> demo -> ask"
               ).font = mw.TITLE_FONT
    note = sheet.cell(
        row=2, column=1,
        value="One row per stage, in the order they run. 'make' commands work on Linux, macOS and "
              "WSL; the scripts\\windows.ps1 equivalents are for native Windows (PowerShell). "
              "The Ollama block below the table explains how the questions are asked with a "
              "local model.")
    note.font, note.alignment = mw.SUBTLE_FONT, mw.WRAP
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    sheet.row_dimensions[2].height = 34
    row = mw.write_header(sheet, ["Step", "Stage", "Command", "Code", "What happens",
                                  "What you can inspect", "Why it is built this way"], 4)
    for entry in PIPELINE:
        row = mw.write_row(sheet, row, list(entry))
        sheet.cell(row=row - 1, column=3).font = mw.MONO_FONT
        sheet.cell(row=row - 1, column=1).font = Font(bold=True)
    row += 1
    sheet.cell(row=row, column=1, value="Asking the questions with Ollama").font = mw.TITLE_FONT
    row += 1
    row = mw.write_header(sheet, ["", "Topic", "Detail", "", "", "", ""], row)
    sheet.freeze_panes = sheet.cell(row=5, column=1)
    for topic, detail in OLLAMA:
        sheet.cell(row=row, column=2, value=topic)
        sheet.cell(row=row, column=3, value=detail)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        for col in range(1, 8):
            cell = sheet.cell(row=row, column=col)
            cell.alignment, cell.border = mw.WRAP, mw.BOX
        sheet.cell(row=row, column=2).font = Font(bold=True)
        if topic.startswith("Status"):
            for col in range(2, 8):
                sheet.cell(row=row, column=col).fill = AMBER
        sheet.row_dimensions[row].height = max(32, 15 * (len(detail) // 150 + detail.count("\n") + 1))
        row += 1


# --------------------------------------------------------------------------
# Tab 4: questions and answers
# --------------------------------------------------------------------------

def sheet_qa(book: Workbook) -> None:
    """Write tab 4 from the recorded LLM run, the reference run and the evaluation.

    :param book: Target workbook.
    :returns: ``None``.
    """
    llm = mw.load_json(config.LLM_RESULTS_JSON, {})
    ref = mw.load_json(config.REFERENCE_RESULTS_JSON, {})
    evaluation = mw.load_json(config.PROJECT_ROOT / "dq" / "results" / "evaluation.json", {})
    verdicts = {v["question_id"]: v for v in evaluation.get("verdicts", [])}
    ref_by_id = _by_id(ref.get("results"))
    llm_by_id = _by_id(llm.get("results"))

    sheet = book.create_sheet("4 Questions & Answers")
    mw.set_widths(sheet, [6, 36, 70, 52, 52, 12, 10, 60, 14, 44, 40])
    sheet.cell(row=1, column=1, value="Questions to SQL: what the LLM generated and what came back"
               ).font = mw.TITLE_FONT
    note = sheet.cell(
        row=2, column=1,
        value=(f"SQL generated by {llm.get('model')} through the {llm.get('provider')} provider, "
               f"executed as database user {llm.get('run_as_role')} (a portfolio holding all "
               f"clients, still filtered by row level security), at git commit "
               f"{llm.get('git_commit')}. The SQL is unedited. It was NOT generated by Ollama; "
               "tab 3 explains how to regenerate it with Ollama. The reference answer is "
               "hand-written SQL (sql/60_reference_queries.sql), never shown to the model. "
               "Every answer was verified against the intent (IR) the model declared before "
               "and after execution (src/sql_ir.py, D029); a guard pass alone is not evidence "
               "of a correct answer. Verdict = value-by-value comparison, tolerance 0.011."))
    note.font, note.alignment = mw.SUBTLE_FONT, mw.WRAP
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=11)
    sheet.row_dimensions[2].height = 48
    row = mw.write_header(sheet, [
        "#", "Question (natural language)", "Generated SQL", "Answer returned by that SQL",
        "Reference answer (hand-written SQL)", "Verdict", "Attempts",
        "Model's stated interpretation and assumptions", "Caveats cited",
        "Declared intent (IR)", "Intent check"], 4)
    row = _qa_rows(sheet, row, config.QUESTION_IDS, llm_by_id, ref_by_id, verdicts)
    tally = evaluation.get("tally", {})
    sheet.cell(row=row + 1, column=2, value="Evaluation tally").font = Font(bold=True)
    sheet.cell(row=row + 1, column=3, value=", ".join(f"{k}: {v}" for k, v in tally.items()))

    extra = {k: mw.load_json(config.SUITES["extra"][f"{k}_json"], {})
             for k in ("llm", "reference", "evaluation")}
    if extra["reference"]:
        row += 3
        title = sheet.cell(row=row, column=1, value=(
            "Additional questions, not in the brief: gender of loans and calendar quarters "
            "(sql/61_reference_extra.sql; assumptions A47, A48). Same provider, same checks."))
        title.font = Font(bold=True, size=12)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        extra_verdicts = {v["question_id"]: v for v in extra["evaluation"].get("verdicts", [])}
        ref_extra = _by_id(extra["reference"].get("results"))
        row = _qa_rows(sheet, row + 1, tuple(ref_extra), _by_id(extra["llm"].get("results")),
                       ref_extra, extra_verdicts)
        tally = extra["evaluation"].get("tally", {})
        sheet.cell(row=row + 1, column=2, value="Evaluation tally").font = Font(bold=True)
        sheet.cell(row=row + 1, column=3, value=", ".join(f"{k}: {v}" for k, v in tally.items()))


def _qa_rows(sheet: Any, row: int, question_ids: tuple[str, ...],
             llm_by_id: dict[str, dict[str, Any]], ref_by_id: dict[str, dict[str, Any]],
             verdicts: dict[str, dict[str, Any]]) -> int:
    """Write one row per question and return the next free row."""
    for qid in question_ids:
        item = llm_by_id.get(qid, {})
        reference = ref_by_id.get(qid, {})
        verdict = verdicts.get(qid, {})
        status = item.get("status", "not_executed")
        sql = (item.get("executed_sql") or item.get("generated_sql") or "NOT EXECUTED").strip()
        assumptions = "\n".join(f"- {a}" for a in item.get("model_assumptions") or [])
        values = [
            qid, item.get("question") or reference.get("question", ""), sql,
            mw.format_rows(item.get("rows") or []) or status,
            mw.format_rows(reference.get("rows") or []),
            verdict.get("verdict", status), item.get("attempts_used", ""),
            (item.get("model_reasoning") or "") + ("\n\n" + assumptions if assumptions else ""),
            ", ".join(item.get("model_caveats") or []) or "-",
            json.dumps(item["ir"], indent=1) if item.get("ir") else "-",
            mw._intent_check(item),  # noqa: SLF001 - same wording as the submission workbook
        ]
        row = mw.write_row(sheet, row, values)
        sheet.cell(row=row - 1, column=3).font = mw.MONO_FONT
        sheet.cell(row=row - 1, column=1).font = Font(bold=True)
        verdict_cell = sheet.cell(row=row - 1, column=6)
        verdict_cell.fill = mw.VERDICT_FILL.get(verdict_cell.value, mw.VERDICT_FILL["not_executed"])
        verdict_cell.font = Font(bold=True)
        lines = max(sql.count("\n"), values[3].count("\n"), values[4].count("\n"),
                    len(values[7]) // 62, 6) + 1
        sheet.row_dimensions[row - 1].height = min(400, 13 * lines)
    return row


def _by_id(results: Any) -> dict[str, dict[str, Any]]:
    """Index a results payload (list or dict) by question id.

    :param results: Contents of a ``results`` key.
    :returns: Mapping from question id to entry.
    """
    if isinstance(results, dict):
        return {key: dict(value, question_id=key) for key, value in results.items()}
    return {item["question_id"]: item for item in results or []}


# --------------------------------------------------------------------------
# Tab 6: design decisions with cost and gain
# --------------------------------------------------------------------------

def _clean(text: str) -> str:
    """Strip markdown noise and join wrapped lines, keeping paragraph breaks.

    :param text: Raw markdown.
    :returns: Plain text.
    """
    text = text.replace("**", "").replace("`", "")
    paragraphs = re.split(r"\n\s*\n", text.strip())
    return "\n".join(" ".join(p.split()) for p in paragraphs if p.strip())


def parse_decisions() -> list[dict[str, str]]:
    """Split DECISIONS.md into one record per decision.

    Gain is the Reason (or "Reason and consequence"); cost is the Cost field
    plus the Consequence. A decision that records neither is left blank rather
    than filled in.

    :returns: Records with id, title, options, choice, gain, cost.
    """
    text = (config.PROJECT_ROOT / "DECISIONS.md").read_text(encoding="utf-8")
    parts = re.split(r"^## (D\d+) - (.*)$", text, flags=re.M)
    records = []
    for i in range(1, len(parts), 3):
        body = re.split(r"^---\s*$", parts[i + 2], flags=re.M)[0]
        fields: list[tuple[str, str]] = []
        pieces = re.split(r"^\*\*([A-Za-z ]+?)\.\*\*", body, flags=re.M)
        for j in range(1, len(pieces), 2):
            fields.append((pieces[j].strip(), _clean(pieces[j + 1])))

        def pick(*names: str) -> str:
            return "\n\n".join(v for k, v in fields if any(k.startswith(n) for n in names))

        options, choice = pick("Options"), pick("Choice")
        letter = re.match(r"^\((\w)\)", choice)
        if letter and len(choice) < 40:
            # "Choice. (b)." names an option; spell it out so the cell stands alone.
            chosen = re.search(rf"\({letter.group(1)}\)(.*?)(?=\s\(\w\)\s|$)", options, flags=re.S)
            if chosen:
                choice = f"({letter.group(1)}){chosen.group(1)}".strip()
        records.append({
            "id": parts[i], "title": _clean(parts[i + 1]),
            "context": pick("Context"),
            "options": options,
            "choice": choice,
            "gain": pick("Reason", "Onboarding", "Result"),
            "cost": pick("Cost", "What it costs", "Consequence"),
            "history": pick("What happened", "How it was caught", "Honest note"),
        })
    return records


def sheet_decisions(book: Workbook) -> int:
    """Write tab 6.

    :param book: Target workbook.
    :returns: Number of decisions.
    """
    records = parse_decisions()
    sheet = book.create_sheet("6 Design decisions")
    mw.set_widths(sheet, [7, 38, 50, 50, 58, 58])
    sheet.cell(row=1, column=1, value="Design decisions: what was chosen, what it gained, what it cost"
               ).font = mw.TITLE_FONT
    note = sheet.cell(
        row=2, column=1,
        value="Taken from DECISIONS.md. 'Gain' is the recorded reason; 'Cost' is the recorded cost "
              "and consequence. Where the log records a defect found after the fact (D008, D010, "
              "D011), the story is included in the gain column. An empty cell means the log "
              "records nothing for it.")
    note.font, note.alignment = mw.SUBTLE_FONT, mw.WRAP
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    sheet.row_dimensions[2].height = 34
    row = mw.write_header(sheet, ["ID", "Decision", "Options considered", "Choice",
                                  "Gain (why this was worth it)", "Cost and consequence"], 4)
    for rec in records:
        gain = rec["gain"] or "Not recorded as a separate reason in DECISIONS.md; see the choice."
        if rec["history"]:
            gain = (rec["history"] + ("\n\n" + gain if gain else "")).strip()
        context = rec["options"] or rec["context"]
        cost = rec["cost"] or ("Recorded together with the reason (see the gain column)."
                               if rec["history"] or "consequence" in rec["gain"].lower()
                               else "Not recorded in DECISIONS.md for this decision.")
        row = mw.write_row(sheet, row, [rec["id"], rec["title"], context, rec["choice"],
                                        gain, cost])
        sheet.cell(row=row - 1, column=1).font = Font(bold=True)
        sheet.cell(row=row - 1, column=5).fill = GREEN
        sheet.cell(row=row - 1, column=6).fill = AMBER
        longest = max(len(rec["options"] or rec["context"]) // 48, len(rec["choice"]) // 48,
                      len(gain) // 56, len(rec["cost"]) // 56, 3)
        sheet.row_dimensions[row - 1].height = min(409, 14 * (longest + 1))
    sheet.auto_filter.ref = f"A4:F{row - 1}"
    return len(records)


# --------------------------------------------------------------------------

def core_counts() -> dict[str, Any]:
    """Read live row counts for the model tab.

    :returns: Mapping from model object name to row count.
    """
    counts: dict[str, Any] = {}
    with db.connect() as conn:
        for name, kind, _, _, _ in build_model.MODEL_OBJECTS:
            if kind.startswith("curated"):
                counts[name] = int(db.fetch_value(conn, f"SELECT count(*) FROM core.{name}") or 0)
            else:
                counts[name] = db.table_count(conn, config.SCHEMA_CORE, name)
    return counts


def main() -> int:
    """Build the workbook.

    :returns: Process exit code.
    """
    book = Workbook()
    book.remove(book.active)
    defects = mw.sheet_defects(book)
    mw.sheet_model(book, {"core_counts": core_counts()})
    sheet_semantic(book)
    sheet_qa(book)
    mw.sheet_assumptions(book)
    decisions = sheet_decisions(book)
    for old, new in (("Defect types", "1 Defects"), ("Model", "2 Model & Explanation"),
                     ("Assumptions", "5 Assumptions")):
        book[old].title = new
    book._sheets.sort(key=lambda s: s.title)  # noqa: SLF001 - tabs are numbered
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    book.save(TARGET)
    print(f"wrote {TARGET.relative_to(ROOT)} -> {' | '.join(book.sheetnames)}")
    print(f"  defect types={defects}  decisions={decisions}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
